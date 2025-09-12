from __future__ import annotations
import os, threading
from pathlib import Path
from typing import List, Dict, Any
from fastapi import APIRouter

router = APIRouter(prefix="/ims", tags=["ims"])

# ---- in-memory index/state ----
IMS_INDEX: List[Dict[str, Any]] = []
IMS_RUNNING = False
IMS_LAST_ERROR: str | None = None
IMS_FILES_SEEN = 0

# ---- config ----
DATA_DIR = Path(os.getenv("DATA_DIR", "data")).resolve()
IMS_DIR = Path(os.getenv("IMS_DIR", DATA_DIR / "source_docs")).resolve()

# ---- lightweight extractors (no asgi import) ----
def _read_txt(fp: Path) -> str:
    try:
        return fp.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""

def _read_docx(fp: Path) -> str:
    try:
        from docx import Document  # python-docx
        d = Document(str(fp))
        parts = []
        for p in getattr(d, "paragraphs", []):
            t = (getattr(p, "text", "") or "").strip()
            if t:
                parts.append(t)
        return "\n".join(parts)
    except Exception:
        return ""

def _read_pdf(fp: Path) -> str:
    try:
        from PyPDF2 import PdfReader
    except Exception:
        return ""
    try:
        pdf = PdfReader(open(fp, "rb"), strict=False)
        return "\n".join((page.extract_text() or "") for page in getattr(pdf, "pages", []))
    except Exception:
        return ""

def _read_xlsx(fp: Path) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filename=str(fp), data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets[:3]:
            out.append(f"# SHEET: {ws.title}")
            count = 0
            for row in ws.iter_rows(values_only=True):
                line = " ".join(str(v) for v in row if v is not None).strip()
                if line:
                    out.append(line)
                    count += 1
                    if count >= 500:
                        break
        return "\n".join(out)
    except Exception:
        return ""

def extract_text_from_file(fp: Path) -> str:
    ext = fp.suffix.lower()
    if ext in (".txt", ".md"): return _read_txt(fp)
    if ext == ".docx": return _read_docx(fp)
    if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"): return _read_xlsx(fp)
    if ext == ".pdf": return _read_pdf(fp)
    return ""

# ---- simple chunker ----
def _chunk(text: str, max_chars: int = 1500) -> List[str]:
    if not text:
        return []
    out, cur = [], ""
    for line in text.splitlines():
        if len(cur) + len(line) + 1 > max_chars:
            if cur.strip():
                out.append(cur.strip())
            cur = line
        else:
            cur = (cur + "\n" if cur else "") + line
    if cur.strip():
        out.append(cur.strip())
    return out[:20]  # cap per-file to keep memory sane

# ---- index builder ----
def _build_index() -> Dict[str, Any]:
    global IMS_INDEX, IMS_RUNNING, IMS_LAST_ERROR, IMS_FILES_SEEN
    IMS_RUNNING = True
    IMS_LAST_ERROR = None
    IMS_FILES_SEEN = 0
    chunks: List[Dict[str, Any]] = []

    try:
        base = IMS_DIR
        if not base.exists():
            raise RuntimeError(f"IMS_DIR not found: {base}")

        for fp in base.rglob("*"):
            if not fp.is_file(): 
                continue
            if any(part.startswith((".__", "._")) or part == "__MACOSX" for part in fp.parts):
                continue
            ext = fp.suffix.lower()
            if ext not in (".txt", ".md", ".docx", ".pdf", ".xlsx", ".xlsm", ".xltx", ".xltm"):
                continue

            rel = str(fp.relative_to(base))
            text = extract_text_from_file(fp)
            if not text.strip():
                continue

            IMS_FILES_SEEN += 1
            for i, seg in enumerate(_chunk(text)):
                chunks.append({"relpath": rel, "i": i, "text": seg})
    except Exception as e:
        IMS_LAST_ERROR = str(e)
    finally:
        IMS_INDEX = chunks
        IMS_RUNNING = False

    return {"files": IMS_FILES_SEEN, "chunks": len(IMS_INDEX), "error": IMS_LAST_ERROR}

# ---- endpoints ----
@router.get("/_ping")
def ims_ping():
    return {"ok": True, "ims_dir": str(IMS_DIR)}

@router.post("/reindex")
def ims_reindex():
    global IMS_RUNNING
    if IMS_RUNNING:
        return {"ok": True, "started": False, "running": True}
    t = threading.Thread(target=_build_index, daemon=True)
    t.start()
    return {"ok": True, "started": True}

@router.get("/status")
def ims_status():
    return {
        "running": IMS_RUNNING,
        "last_error": IMS_LAST_ERROR,
        "files_seen": IMS_FILES_SEEN,
        "chunks": len(IMS_INDEX),
        "ims_dir": str(IMS_DIR),
    }
